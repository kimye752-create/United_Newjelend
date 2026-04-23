#!/usr/bin/env python3
"""뉴질랜드 시장 분석 보고서 DOCX 생성기.

SG 템플릿 양식 (나눔고딕) 기준으로:
  - build_cover_docx   → 표지 (NZ_00_표지)
  - build_p1_docx      → 시장보고서 (NZ_01)
  - build_p2_docx      → 수출가격전략 보고서 (NZ_02)
  - build_p3_docx      → 바이어 후보 리스트 (NZ_03)
  - build_final_docx   → 최종 합본 (표지 + P2 + P3 + P1)

사용:
  from report_generator_docx import build_p1_docx, build_p2_docx, build_p3_docx, build_final_docx
  path = build_p1_docx(data, Path("reports/nz_p1_xxx.docx"))
"""
from __future__ import annotations

import copy
import io
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor, Inches

# ── 색상 팔레트 ──────────────────────────────────────────────────────────────
C_NAVY   = RGBColor(0x17, 0x3F, 0x78)   # 진한 네이비
C_DARK   = RGBColor(0x1F, 0x2D, 0x3D)   # 다크 타이틀
C_BODY   = RGBColor(0x33, 0x33, 0x33)   # 본문
C_MUTED  = RGBColor(0x88, 0x88, 0x88)   # 부연
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT  = RGBColor(0xF2, 0xF5, 0xFA)   # 표 헤더 배경
C_BORDER = RGBColor(0xCC, 0xCC, 0xCC)

# ── 폰트 ─────────────────────────────────────────────────────────────────────
FONT_KO  = "나눔고딕"
FONT_EN  = "나눔고딕"  # 동일 폰트 통일

TODAY_STR = date.today().strftime("%Y-%m-%d")


# ─────────────────────────────────────────────────────────────────────────────
#  헬퍼 함수
# ─────────────────────────────────────────────────────────────────────────────

def _set_font(run, size_pt: float, bold: bool = False, color: RGBColor | None = None,
              italic: bool = False):
    run.font.name = FONT_KO
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    # 동아시아 폰트도 동일하게
    rPr = run._r
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_KO)


def _para_spacing(para, before_pt: float = 0, after_pt: float = 0, line_pt: float | None = None):
    pf = para.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after  = Pt(after_pt)
    if line_pt:
        from docx.shared import Pt as _Pt
        pf.line_spacing = _Pt(line_pt)


def _add_heading(doc: Document, text: str, level: int = 1,
                 size_pt: float = 13, color: RGBColor = C_NAVY,
                 before_pt: float = 12, after_pt: float = 4) -> Any:
    para = doc.add_paragraph()
    _para_spacing(para, before_pt, after_pt)
    run = para.add_run(text)
    _set_font(run, size_pt, bold=True, color=color)
    return para


def _add_body(doc: Document, text: str, size_pt: float = 9.5,
              before_pt: float = 2, after_pt: float = 2) -> Any:
    text = text or "—"
    para = doc.add_paragraph()
    _para_spacing(para, before_pt, after_pt)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    _set_font(run, size_pt, color=C_BODY)
    return para


def _add_bullet(doc: Document, text: str, size_pt: float = 9.5) -> Any:
    para = doc.add_paragraph(style="List Bullet")
    _para_spacing(para, 0, 1)
    run = para.add_run(text)
    _set_font(run, size_pt, color=C_BODY)
    return para


def _shade_cell(cell, fill_hex: str):
    """셀 배경색 지정."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  fill_hex)
    tcPr.append(shd)


def _cell_text(cell, text: str, size_pt: float = 9, bold: bool = False,
               color: RGBColor = C_BODY, align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    _para_spacing(para, 1, 1)
    run = para.add_run(str(text) if text is not None else "—")
    _set_font(run, size_pt, bold=bold, color=color)


def _add_kv_table(doc: Document, rows: list[tuple[str, str]],
                  col_widths_cm: tuple[float, float] = (4.5, 12.0)):
    """키-값 2열 테이블."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    w0 = Cm(col_widths_cm[0])
    w1 = Cm(col_widths_cm[1])
    for i, (k, v) in enumerate(rows):
        r = table.rows[i]
        r.cells[0].width = w0
        r.cells[1].width = w1
        _shade_cell(r.cells[0], "EBF0FA")
        _cell_text(r.cells[0], k, size_pt=9, bold=True, color=C_NAVY)
        _cell_text(r.cells[1], v or "—", size_pt=9)
    _para_spacing(doc.add_paragraph(), 3, 3)
    return table


def _add_horizontal_rule(doc: Document):
    """가로 구분선."""
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C0C8D8")
    pBdr.append(bottom)
    pPr.append(pBdr)
    _para_spacing(para, 4, 4)


def _page_break(doc: Document):
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(__import__("docx.oxml.ns", fromlist=["qn"])  # noqa
                  and __import__("docx", fromlist=["oxml"]).oxml.OxmlElement("w:br"))
    # Simpler approach:
    para = doc.add_paragraph()
    para.runs  # ensure exists
    para.clear()
    pPr = para._p.get_or_add_pPr()
    from docx.oxml import OxmlElement as _E
    pgBr = _E("w:r")
    br = _E("w:br")
    br.set(qn("w:type"), "page")
    pgBr.append(br)
    para._p.append(pgBr)


def _set_page_margins(doc: Document,
                      top: float = 2.5, bottom: float = 2.5,
                      left: float = 2.5, right: float = 2.0):
    """페이지 여백 (cm)."""
    for section in doc.sections:
        section.top_margin    = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin   = Cm(left)
        section.right_margin  = Cm(right)


# ─────────────────────────────────────────────────────────────────────────────
#  표지 (Cover)
# ─────────────────────────────────────────────────────────────────────────────

def build_cover_docx(data: dict, out_path: Path | None = None) -> Document:
    """
    data keys:
      country_ko    : str  (예: "뉴질랜드")
      company       : str  (예: "한국유나이티드제약")
      report_date   : str  (예: "2026-04-22")
      subtitle      : str  (예: "수출가격 전략 - 바이어 후보 리스트 - 시장분석")
    """
    doc = Document()
    _set_page_margins(doc)

    # 상단 여백
    for _ in range(8):
        doc.add_paragraph()

    # 메인 타이틀
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(para, 0, 6)
    run = para.add_run(f"{data.get('country_ko', '뉴질랜드')} 진출 전략 보고서")
    _set_font(run, 26, bold=True, color=C_DARK)

    # 회사명
    para2 = doc.add_paragraph()
    para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(para2, 4, 4)
    run2 = para2.add_run(data.get("company", "한국유나이티드제약(주)"))
    _set_font(run2, 14, color=C_NAVY)

    # 날짜
    para3 = doc.add_paragraph()
    para3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(para3, 2, 2)
    run3 = para3.add_run(data.get("report_date", TODAY_STR))
    _set_font(run3, 12, color=C_MUTED)

    # 구분선
    for _ in range(2):
        doc.add_paragraph()
    _add_horizontal_rule(doc)

    # 부제목
    para4 = doc.add_paragraph()
    para4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para_spacing(para4, 6, 0)
    subtitle = data.get("subtitle", "수출가격 전략 - 바이어 후보 리스트 - 시장분석")
    run4 = para4.add_run(subtitle)
    _set_font(run4, 11, color=C_NAVY)

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
    return doc


# ─────────────────────────────────────────────────────────────────────────────
#  P1 — 시장보고서
# ─────────────────────────────────────────────────────────────────────────────

def build_p1_docx(data: dict, out_path: Path | None = None) -> Document:
    """
    data keys (분석 결과 딕셔너리):
      trade_name, inn_name, hs_code, country_ko, report_date
      macro_rows        : list[tuple[str,str]]   # 거시환경 지표
      macro_summary     : str
      regulatory        : dict { registration, channel, tariff }
      ref_prices        : list[dict { label, value, note }]
      ref_price_summary : str
      risks             : dict { review_period, competition, formulary }
      refs              : list[dict { title, summary, url }]
      db_sources        : list[str]
    """
    doc = Document()
    _set_page_margins(doc)

    country_ko  = data.get("country_ko",  "뉴질랜드")
    trade_name  = data.get("trade_name",  "")
    inn_name    = data.get("inn_name",    "")
    hs_code     = data.get("hs_code",     "3004.90")
    rpt_date    = data.get("report_date", TODAY_STR)

    # ── 헤더 타이틀
    para = doc.add_paragraph()
    _para_spacing(para, 0, 4)
    run = para.add_run(f"{country_ko} 시장보고서 — {trade_name}")
    _set_font(run, 16, bold=True, color=C_DARK)

    # ── 서브헤더
    para2 = doc.add_paragraph()
    _para_spacing(para2, 0, 10)
    sub_text = f"{trade_name} ({inn_name})  |  HS CODE: {hs_code}  |  {country_ko}  |  {rpt_date}"
    run2 = para2.add_run(sub_text)
    _set_font(run2, 9, color=C_MUTED)

    _add_horizontal_rule(doc)

    # ── 1. 의료 거시환경
    _add_heading(doc, "1. 의료 거시환경 파악", size_pt=12)

    macro_rows = data.get("macro_rows") or []
    if macro_rows:
        _add_kv_table(doc, [(k, v) for k, v in macro_rows], col_widths_cm=(5.0, 11.5))

    summary = data.get("macro_summary", "")
    if summary:
        _add_body(doc, summary)

    _add_horizontal_rule(doc)

    # ── 2. 무역/규제 환경
    _add_heading(doc, "2. 무역/규제 환경", size_pt=12)

    reg = data.get("regulatory") or {}
    sections_reg = [
        ("▸ 등록 현황",      reg.get("registration", "—")),
        ("▸ 진입 채널 권고", reg.get("channel",       "—")),
        ("▸ 관세 및 무역",   reg.get("tariff",        "—")),
    ]
    for label, content in sections_reg:
        p = doc.add_paragraph()
        _para_spacing(p, 6, 2)
        r = p.add_run(label)
        _set_font(r, 10, bold=True, color=C_NAVY)
        _add_body(doc, content)

    _add_horizontal_rule(doc)

    # ── 3. 참고 가격
    _add_heading(doc, "3. 참고 가격", size_pt=12)

    ref_prices = data.get("ref_prices") or []
    for rp in ref_prices:
        rows = [
            ("항목",   rp.get("label", "—")),
            ("가격",   rp.get("value", "—")),
        ]
        if rp.get("note"):
            rows.append(("비고", rp["note"]))
        _add_kv_table(doc, rows)

    price_summary = data.get("ref_price_summary", "")
    if price_summary:
        _add_body(doc, price_summary)

    _add_horizontal_rule(doc)

    # ── 4. 리스크 / 조건
    _add_heading(doc, "4. 리스크 / 조건", size_pt=12)

    risks = data.get("risks") or {}
    risk_sections = [
        ("▸ 규제 심사 소요 기간", risks.get("review_period", "—")),
        ("▸ 경쟁 강도",           risks.get("competition",   "—")),
        ("▸ 포뮬러리 등재",       risks.get("formulary",     "—")),
    ]
    for label, content in risk_sections:
        p = doc.add_paragraph()
        _para_spacing(p, 6, 2)
        r = p.add_run(label)
        _set_font(r, 10, bold=True, color=C_NAVY)
        _add_body(doc, content)

    _add_horizontal_rule(doc)

    # ── 5. 근거 및 출처
    _add_heading(doc, "5. 근거 및 출처", size_pt=12)

    # 5-1. Perplexity 논문
    p = doc.add_paragraph()
    _para_spacing(p, 6, 2)
    r = p.add_run("▸ 5-1. Perplexity 추천 논문")
    _set_font(r, 10, bold=True, color=C_NAVY)

    refs = data.get("refs") or []
    if refs:
        for idx, ref in enumerate(refs, 1):
            p_no = doc.add_paragraph()
            _para_spacing(p_no, 4, 0)
            rn = p_no.add_run(f"No.{idx}  {ref.get('title', '—')}")
            _set_font(rn, 9.5, bold=True, color=C_BODY)

            summary_txt = ref.get("summary") or ref.get("description") or ""
            if summary_txt:
                _add_body(doc, summary_txt, size_pt=9)

            url = ref.get("url") or ref.get("link") or ""
            if url:
                p_url = doc.add_paragraph()
                _para_spacing(p_url, 1, 4)
                ru = p_url.add_run(f"출처: {url}")
                _set_font(ru, 8.5, color=C_MUTED, italic=True)
    else:
        _add_body(doc, "—")

    # 5-2. DB/기관
    p2 = doc.add_paragraph()
    _para_spacing(p2, 10, 2)
    r2 = p2.add_run("▸ 5-2. 사용된 DB/기관")
    _set_font(r2, 10, bold=True, color=C_NAVY)

    db_sources = data.get("db_sources") or []
    if db_sources:
        for src in db_sources:
            _add_bullet(doc, src)
    else:
        _add_body(doc, "—")

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
    return doc


# ─────────────────────────────────────────────────────────────────────────────
#  P2 — 수출가격전략 보고서
# ─────────────────────────────────────────────────────────────────────────────

def build_p2_docx(data: dict, out_path: Path | None = None) -> Document:
    """
    data keys:
      trade_name, inn_name, country_ko, report_date
      macro_text        : str   (거시 시장 서술)
      base_price_nzd    : float
      base_price_usd    : float
      base_price_krw    : float
      price_method      : str
      market_type       : str   ("공공 / 민간")
      competitors       : list[dict { name, product, spec, price_text }]
      scenarios_public  : list[dict { label, price_nzd, price_usd, reason, fob_formula }]
      scenarios_private : list[dict { label, price_nzd, price_usd, reason, fob_formula }]
      fx_snapshot       : dict { nzd_krw, nzd_usd, ... }
    """
    doc = Document()
    _set_page_margins(doc)

    country_ko = data.get("country_ko",  "뉴질랜드")
    trade_name = data.get("trade_name",  "")
    inn_name   = data.get("inn_name",    "")
    rpt_date   = data.get("report_date", TODAY_STR)
    fx         = data.get("fx_snapshot") or {}
    nzd_krw    = fx.get("nzd_krw",  830.0)
    nzd_usd    = fx.get("nzd_usd",  0.595)

    def _nzd_to_usd(nzd: float | None) -> str:
        if nzd is None:
            return "—"
        return f"USD {nzd * nzd_usd:,.2f}"

    def _nzd_to_krw(nzd: float | None) -> str:
        if nzd is None:
            return "—"
        return f"KRW {nzd * nzd_krw:,.0f}"

    # ── 헤더 타이틀
    para = doc.add_paragraph()
    _para_spacing(para, 0, 4)
    run = para.add_run(f"{country_ko} 수출 가격 전략 보고서 — {trade_name}")
    _set_font(run, 16, bold=True, color=C_DARK)

    para2 = doc.add_paragraph()
    _para_spacing(para2, 0, 10)
    sub = f"{trade_name} ({inn_name})  | {rpt_date}"
    run2 = para2.add_run(sub)
    _set_font(run2, 9, color=C_MUTED)

    _add_horizontal_rule(doc)

    # ── 1. 거시 시장
    _add_heading(doc, f"1. {country_ko} 거시 시장", size_pt=12)
    _add_body(doc, data.get("macro_text", "—"))

    _add_horizontal_rule(doc)

    # ── 2. 단가 (시장 기준가)
    _add_heading(doc, f"2. {trade_name} 단가 (시장 기준가)", size_pt=12)

    bp_nzd = data.get("base_price_nzd")
    bp_usd = data.get("base_price_usd") or (bp_nzd * nzd_usd if bp_nzd else None)
    bp_krw = data.get("base_price_krw") or (bp_nzd * nzd_krw if bp_nzd else None)

    base_str = "—"
    if bp_nzd:
        base_str = f"USD {bp_usd:,.2f}  /  NZD {bp_nzd:,.2f}  /  KRW {bp_krw:,.0f}" if bp_usd else f"NZD {bp_nzd:,.2f}"

    _add_kv_table(doc, [
        ("기준 가격",  base_str),
        ("산정 방식",  data.get("price_method", "AI 분석 (Claude Haiku)")),
        ("시장 구분",  data.get("market_type",  "공공 / 민간")),
    ])

    _add_horizontal_rule(doc)

    # ── 3. 거래처 참고 가격
    _add_heading(doc, "3. 거래처 참고 가격", size_pt=12)

    competitors = data.get("competitors") or []
    if competitors:
        for comp in competitors:
            _add_kv_table(doc, [
                ("업체명",    comp.get("name",    "—")),
                ("제품명",    comp.get("product", "—")),
                ("성분·함량", comp.get("spec",    "—")),
                ("시장가",    comp.get("price_text", "—")),
            ])
    else:
        _add_body(doc, "— 경쟁사 데이터 없음")

    _add_horizontal_rule(doc)

    # ── 4. 가격 시나리오
    _add_heading(doc, "4. 가격 시나리오", size_pt=12)

    def _render_scenarios(label: str, scenarios: list[dict]):
        ph = doc.add_paragraph()
        _para_spacing(ph, 6, 4)
        rh = ph.add_run(label)
        _set_font(rh, 10, bold=True, color=C_NAVY)

        if not scenarios:
            _add_body(doc, "—")
            return

        for sc in scenarios:
            sc_label = sc.get("label", "")
            price_nzd = sc.get("price_nzd") or sc.get("price") or sc.get("price_sgd")
            price_usd = sc.get("price_usd")
            if price_nzd and not price_usd:
                price_usd = price_nzd * nzd_usd
            price_krw = price_nzd * nzd_krw if price_nzd else None

            if price_nzd:
                price_display = f"NZD {price_nzd:,.2f}"
                if price_usd:
                    price_display = f"USD {price_usd:,.2f}  /  {price_display}"
                if price_krw:
                    price_display += f"  /  KRW {price_krw:,.0f}"
            else:
                price_display = "—"

            # 레이블 + 가격
            p_sc = doc.add_paragraph()
            _para_spacing(p_sc, 6, 2)
            r_sc = p_sc.add_run(f"[{sc_label}]  {price_display}")
            _set_font(r_sc, 10, bold=True, color=C_BODY)

            # 근거
            reason = sc.get("reason", "")
            if reason:
                p_rs = doc.add_paragraph()
                _para_spacing(p_rs, 1, 2)
                r_rs = p_rs.add_run("근거  ")
                _set_font(r_rs, 9, bold=True, color=C_NAVY)
                r_rs2 = p_rs.add_run(reason)
                _set_font(r_rs2, 9, color=C_BODY)

            # FOB 역산식
            fob = sc.get("fob_formula", "")
            if fob:
                p_fob = doc.add_paragraph()
                _para_spacing(p_fob, 1, 6)
                r_fob = p_fob.add_run("FOB 수출가 역산식  ")
                _set_font(r_fob, 9, bold=True, color=C_NAVY)
                r_fob2 = p_fob.add_run(fob)
                _set_font(r_fob2, 9, color=C_BODY)

    scenarios_pub  = data.get("scenarios_public")  or data.get("scenarios") or []
    scenarios_priv = data.get("scenarios_private") or []

    _render_scenarios(
        f"▸ 4-1. 공공 시장  (데이터 소스: PHARMAC 급여가, GETS 조달가 참고)",
        scenarios_pub
    )
    _render_scenarios(
        f"▸ 4-2. 민간 시장  (데이터 소스: 민간 병원·약국 공급가 참고)",
        scenarios_priv
    )

    _add_horizontal_rule(doc)

    # ── 면책 문구
    para_disc = doc.add_paragraph()
    _para_spacing(para_disc, 6, 0)
    run_disc = para_disc.add_run(
        "※ 본 산출 결과는 AI 분석에 기반한 추정치이므로, "
        "최종 의사결정 전 반드시 담당자의 검토 및 확인이 필요합니다."
    )
    _set_font(run_disc, 8.5, color=C_MUTED, italic=True)

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
    return doc


# ─────────────────────────────────────────────────────────────────────────────
#  P3 — 바이어 후보 리스트
# ─────────────────────────────────────────────────────────────────────────────

def build_p3_docx(data: dict, out_path: Path | None = None) -> Document:
    """
    data keys:
      trade_name, country_ko, report_date
      buyers : list[dict]
        each buyer:
          rank, name, country_en, category, email, website,
          address, phone, company_size, products,
          overview        : str
          reasons         : list[str] (5개, ①~⑤ 접두 없어도 됨)
          reason_labels   : list[str] (선택, 기본 ["매출규모","파이프라인","제조소 보유","수입 경험","약국 체인 운영"])
    """
    doc = Document()
    _set_page_margins(doc)

    country_ko = data.get("country_ko",  "뉴질랜드")
    trade_name = data.get("trade_name",  "")
    rpt_date   = data.get("report_date", TODAY_STR)
    buyers     = data.get("buyers") or []

    REASON_LABELS_DEFAULT = ["매출 규모", "파이프라인", "제조소 보유", "수입 경험", "약국 체인 운영"]
    RANK_NUMS = ["①", "②", "③", "④", "⑤"]

    # ── 헤더
    para = doc.add_paragraph()
    _para_spacing(para, 0, 4)
    run = para.add_run(f"{country_ko} 바이어 후보 리스트 — {trade_name}")
    _set_font(run, 16, bold=True, color=C_DARK)

    para2 = doc.add_paragraph()
    _para_spacing(para2, 0, 6)
    run2 = para2.add_run(f"{country_ko}  |  {rpt_date}")
    _set_font(run2, 9, color=C_MUTED)

    # 주의사항
    p_note = doc.add_paragraph()
    _para_spacing(p_note, 0, 8)
    r_note = p_note.add_run(
        "※ 아래 바이어 후보는 CPHI 등록 및 Perplexity 웹 분석을 통해 도출되었으며, "
        "개별 기업의 뉴질랜드 진출 현황 및 제품 연관성은 추가 실사가 필요합니다."
    )
    _set_font(r_note, 8.5, color=C_MUTED, italic=True)

    _add_horizontal_rule(doc)

    # ── 1. 바이어 후보 리스트 (전체)
    _add_heading(doc, f"1. 바이어 후보 리스트 (전체 {len(buyers)}개사)", size_pt=12)

    if buyers:
        # 리스트 표: 번호 | 업체명 | 국가 | 카테고리 | 이메일
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for cell, txt in zip(hdr, ["No.", "업체명", "국가", "카테고리", "이메일"]):
            _shade_cell(cell, "173F78")
            _cell_text(cell, txt, size_pt=8.5, bold=True, color=C_WHITE)

        for i, b in enumerate(buyers, 1):
            row = table.add_row().cells
            _cell_text(row[0], str(i),                         size_pt=8.5)
            _cell_text(row[1], b.get("name",        "—"),      size_pt=8.5, bold=True)
            _cell_text(row[2], b.get("country_en",  "—"),      size_pt=8.5)
            _cell_text(row[3], b.get("category",    "—"),      size_pt=8.5)
            _cell_text(row[4], b.get("email",       "—"),      size_pt=8.5)

    doc.add_paragraph()
    _add_horizontal_rule(doc)

    # ── 2. 우선 접촉 바이어 상세 정보 (TOP 10 전원)
    _add_heading(doc, "2. 우선 접촉 바이어 상세 정보 (상위 10개사)", size_pt=12)

    p_note2 = doc.add_paragraph()
    _para_spacing(p_note2, 0, 8)
    r_note2 = p_note2.add_run(
        f"※ 하기 기업은 {trade_name}의 성분 연관성, NZ 지역 네트워크, "
        "진출 가능성을 종합 평가하여 선정하였습니다."
    )
    _set_font(r_note2, 8.5, color=C_MUTED, italic=True)

    detail_buyers = buyers[:10]  # TOP 10
    for b in detail_buyers:
        rank     = b.get("rank", buyers.index(b) + 1 if b in buyers else "?")
        name     = b.get("name",     "—")
        cat      = b.get("category", "—")
        country  = b.get("country_en", "—")

        # 업체 제목
        p_title = doc.add_paragraph()
        _para_spacing(p_title, 12, 2)
        r_rank = p_title.add_run(f"{rank}. ")
        _set_font(r_rank, 11, bold=True, color=C_NAVY)
        r_name = p_title.add_run(f"{name}")
        _set_font(r_name, 11, bold=True, color=C_DARK)
        r_meta = p_title.add_run(f"  |  {country} · {cat}")
        _set_font(r_meta, 9, color=C_MUTED)

        # 기업 개요
        p_ov_lbl = doc.add_paragraph()
        _para_spacing(p_ov_lbl, 4, 1)
        r_ov_lbl = p_ov_lbl.add_run("▸ 기업 개요")
        _set_font(r_ov_lbl, 9.5, bold=True, color=C_NAVY)

        overview = b.get("overview") or b.get("summary") or "—"
        _add_body(doc, overview, size_pt=9)

        # 추천 이유 (5가지)
        p_re_lbl = doc.add_paragraph()
        _para_spacing(p_re_lbl, 6, 1)
        r_re_lbl = p_re_lbl.add_run("▸ 추천 이유")
        _set_font(r_re_lbl, 9.5, bold=True, color=C_NAVY)

        reasons = b.get("reasons") or []
        reason_labels = b.get("reason_labels") or REASON_LABELS_DEFAULT
        for ri, (rnum, rlbl) in enumerate(zip(RANK_NUMS, reason_labels)):
            reason_text = reasons[ri] if ri < len(reasons) else "—"
            p_r = doc.add_paragraph()
            _para_spacing(p_r, 2, 1)
            r_lbl = p_r.add_run(f"{rnum} {rlbl}  ")
            _set_font(r_lbl, 9, bold=True, color=C_BODY)
            r_txt = p_r.add_run(reason_text)
            _set_font(r_txt, 9, color=C_BODY)

        # 연락처 표
        contact_rows: list[tuple[str, str]] = []
        for field, key in [
            ("주소",     "address"),
            ("전화",     "phone"),
            ("이메일",   "email"),
            ("홈페이지", "website"),
            ("기업 규모","company_size"),
            ("등록 제품","products"),
        ]:
            val = b.get(key, "")
            contact_rows.append((field, val or "—"))

        _add_kv_table(doc, contact_rows, col_widths_cm=(3.5, 13.0))

        # 출처
        p_src = doc.add_paragraph()
        _para_spacing(p_src, 0, 2)
        r_src = p_src.add_run("※ 출처: Perplexity 분석")
        _set_font(r_src, 8, color=C_MUTED, italic=True)

        _add_horizontal_rule(doc)

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
    return doc


# ─────────────────────────────────────────────────────────────────────────────
#  Final — 최종 합본 (표지 + P2 + P3 + P1)
# ─────────────────────────────────────────────────────────────────────────────

def _append_doc(base: Document, addition: Document):
    """addition의 body 요소들을 base에 이어 붙인다 (페이지 브레이크 포함)."""
    from docx.oxml import OxmlElement as _E
    from docx.oxml.ns import qn as _qn

    # 페이지 브레이크 삽입
    br_para = _E("w:p")
    br_r    = _E("w:r")
    br_el   = _E("w:br")
    br_el.set(_qn("w:type"), "page")
    br_r.append(br_el)
    br_para.append(br_r)
    base.element.body.append(br_para)

    # addition의 body 요소 복사 (마지막 sectPr 제외)
    for elem in addition.element.body:
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "sectPr":
            continue
        base.element.body.append(copy.deepcopy(elem))


def build_final_docx(
    cover_data:  dict,
    p2_data:     dict,
    p3_data:     dict,
    p1_data:     dict,
    out_path:    Path | None = None,
) -> Document:
    """최종 합본 DOCX 생성: 표지 + P2(가격전략) + P3(바이어) + P1(시장보고서)."""

    # 각 파트 생성
    cover_doc = build_cover_docx(cover_data)
    p2_doc    = build_p2_docx(p2_data)
    p3_doc    = build_p3_docx(p3_data)
    p1_doc    = build_p1_docx(p1_data)

    # 표지를 베이스로 병합
    final = cover_doc
    _append_doc(final, p2_doc)
    _append_doc(final, p3_doc)
    _append_doc(final, p1_doc)

    if out_path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        final.save(str(out_path))
    return final


# ─────────────────────────────────────────────────────────────────────────────
#  기존 render_pdf 호환 래퍼 (서버에서 호출하던 함수명 유지)
# ─────────────────────────────────────────────────────────────────────────────

def render_p1_docx(report: dict, out_path: Path) -> None:
    """서버에서 호출: P1 시장보고서 DOCX 저장."""
    build_p1_docx(report, out_path)


def render_p2_docx_report(p2_data: dict, out_path: Path) -> None:
    """서버에서 호출: P2 가격전략 DOCX 저장."""
    build_p2_docx(p2_data, out_path)


def render_p3_docx_report(buyers: list, trade_name: str, out_path: Path,
                           country_ko: str = "뉴질랜드",
                           report_date: str | None = None) -> None:
    """서버에서 호출: P3 바이어 DOCX 저장."""
    data = {
        "trade_name":  trade_name,
        "country_ko":  country_ko,
        "report_date": report_date or TODAY_STR,
        "buyers":      buyers,
    }
    build_p3_docx(data, out_path)


def render_final_docx(cover_data: dict, p2_data: dict,
                       p3_data: dict, p1_data: dict, out_path: Path) -> None:
    """서버에서 호출: 최종 합본 DOCX 저장."""
    build_final_docx(cover_data, p2_data, p3_data, p1_data, out_path)
